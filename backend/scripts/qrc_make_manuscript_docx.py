"""Render the QRC manuscript markdown to a Word .docx (python-docx).

A focused markdown->docx converter for QRC_Manuscript_Full.md: headings, paragraphs,
pipe tables, images (![](figures/...)), **bold**, `code`, blockquotes, bullet lists,
and fenced code blocks. Not a general converter — just enough for this manuscript.

Run: python backend/scripts/qrc_make_manuscript_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS = Path(__file__).resolve().parents[2] / "docs" / "QRC"
SRC = DOCS / "QRC_Manuscript_Full.md"
OUT = DOCS / "QRC_Manuscript.docx"

_IMG = re.compile(r"^!\[.*?\]\((.*?)\)\s*$")
_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _add_runs(par, text):
    """Add text to a paragraph honoring **bold** and `code` inline spans."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        else:
            par.add_run(part)


def _emit_table(doc, rows):
    """rows: list of markdown rows (incl. the |---| separator at index 1)."""
    def cells(line):
        return [c.strip() for c in line.strip().strip("|").split("|")]
    header = cells(rows[0])
    body = [cells(r) for r in rows[2:]]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(header):
        _add_runs(t.rows[0].cells[j].paragraphs[0], h)
    for row in body:
        rc = t.add_row().cells
        for j in range(len(header)):
            _add_runs(rc[j].paragraphs[0], row[j] if j < len(row) else "")
    doc.add_paragraph()


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)

    i = 0
    in_code = False
    code_buf = []
    while i < len(lines):
        ln = lines[i]

        if ln.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf)); r.font.name = "Consolas"; r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_buf = []; in_code = False
            else:
                in_code = True
            i += 1; continue
        if in_code:
            code_buf.append(ln); i += 1; continue

        m = _IMG.match(ln)
        if m:
            img = (DOCS / m.group(1)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue

        if ln.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            _emit_table(doc, block); continue

        if ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            txt = ln.lstrip("#").strip()
            doc.add_heading(txt, level=min(lvl, 4) - (1 if lvl == 1 else 0) if lvl > 1 else 0)
            i += 1; continue

        if ln.strip() == "---":
            i += 1; continue

        if ln.strip().startswith("> "):
            p = doc.add_paragraph(); p.style = "Intense Quote"
            _add_runs(p, ln.strip()[2:]); i += 1; continue

        if ln.strip().startswith("- ") or ln.strip().startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, ln.strip()[2:]); i += 1; continue

        if ln.strip() == "":
            i += 1; continue

        p = doc.add_paragraph()
        _add_runs(p, ln)
        i += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT}  ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
